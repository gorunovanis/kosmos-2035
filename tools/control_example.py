"""Independent manual control example in XLSX (Критерий 5): the annual algebra of the chosen plan for
2037 (BASE) and 2039 (MANDATORY_STRESS, prepared plan) written as Excel formulas from the inputs, next
to the engine's numbers. Opening the workbook in Excel/LibreOffice recomputes the formulas.

Usage: python tools/control_example.py  -> docs/control_example.xlsx"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from kosmo.api import CASE_DIR, CONFIG_DIR, all_scenarios, load_assumptions  # noqa: E402
from kosmo.case import load_case  # noqa: E402
from kosmo.engine import run  # noqa: E402
from kosmo.plan import Plan  # noqa: E402

CHOSEN = "S1_new24"


def sheet_for_year(wb, title, case, plan, scen, params, year):
    res = run(case, plan, scen, params, keep_daily=False)
    row = next(r for r in res.yearly if r["year"] == year)
    prev = next((r for r in res.yearly if r["year"] == year - 1), None)
    src = {r["source_id"]: r for r in res.source_year if r["year"] == year}
    from kosmo.scenario import effective
    eff = effective(case, scen, case.years)
    ws = wb.create_sheet(title)
    bold = Font(bold=True)
    ws.append(["Независимый контрольный пример", f"план {plan.plan_id}, сценарий {scen.scenario_id}, год {year}"])
    ws["A1"].font = bold
    ws.append(["Формулы в столбце C считаются Excel из входов столбца B; столбец D — числа ядра из выгрузки; E — разница."])
    ws.append([])
    ws.append(["Вход", "Значение", "Формула / расчёт", "Ядро", "Разница"])
    for c in "ABCDE":
        ws[f"{c}4"].font = bold
    r0 = 5
    inputs = [
        ("Спрос total, т", row["demand_total_t"]),
        ("Спрос critical, т", row["demand_critical_t"]),
        ("Запас на 1 января, т (из ядра, конец предыдущего года)", row["opening_inventory_t"]),
        ("Коэффициент потерь (ZBO 1,2 % / база 4,5 %)", 0.012 if (res.meta.get("zbo_commissioning") and res.meta["zbo_commissioning"] <= f"{year}-01-01") else 0.045),
        ("Дней в году", 365),
        ("Резервных дней", params["reserve_days"]),
        ("Ставка дисконта", params["discount_rate"]),
        ("Год приведения", params["discount_base_year"]),
        ("Ставка хранения, млн/т-год", case.storages["BASE"].holding_cost),
        ("OPEX ZBO, млн/год", case.storages["ZBO"].opex),
    ]
    for name, val in inputs:
        ws.append([name, val])
    r_dem, r_crit, r_open, r_loss, r_days, r_rdays, r_rate, r_t0, r_hold, r_opex = range(r0, r0 + 10)
    ws.append([])
    ws.append(["Канал", "Заказано, т", "Цена, млн/т", "Резерв, т/год", "Доля периода", "TOP доля", "Тариф резерва", "Оплачиваемый объём = max(заказ, TOP×резерв×доля)", "Закупка = цена × оплачиваемый", "Резервирование = тариф × резерв × доля", "Ядро: закупка", "Ядро: резерв"])
    hdr = ws.max_row
    for c in "ABCDEFGHIJKL":
        ws[f"{c}{hdr}"].font = bold
    first = ws.max_row + 1
    for sid, s in case.sources.items():
        sr = src.get(sid)
        if sr is None:
            continue
        rr = ws.max_row + 1
        ws.append([s.name, sr["scheduled_t"], sr["unit_price_mln_per_t"], sr["reserved_t_per_year"], sr["contract_period_fraction"], s.top_share, sr["reservation_rate_mln_per_t_year"],
                   f"=MAX(B{rr},F{rr}*D{rr}*E{rr})", f"=C{rr}*H{rr}", f"=G{rr}*D{rr}*E{rr}", sr["procurement_mln"], sr["reservation_mln"]])
    last = ws.max_row
    ws.append([])
    ws.append(["Показатель", "", "Формула", "Ядро", "Разница"])
    for c in "ACDE":
        ws[f"{c}{ws.max_row}"].font = bold
    calc = [
        ("Заказано всего, т", f"=SUM(B{first}:B{last})", row["scheduled_t"]),
        ("Доставлено, т (сценарные доли поставки учтены в ядре; для земных каналов = заказу)", f"=SUM(B{first}:B{last})", row["delivered_t"]),
        ("Потери = доставлено × коэффициент, т", f"=C{{deliv}}*B{r_loss}", row["losses_t"]),
        ("Резерв 45 дн = спрос × дни / 365, т", f"=B{r_dem}*B{r_rdays}/B{r_days}", row["required_opening_reserve_t"]),
        ("Проверка резерва: запас на 1 января − резерв, т (≥ 0)", f"=B{r_open}-C{{reserve}}", row["opening_inventory_t"] - row["required_opening_reserve_t"]),
        ("Выдано = min(спрос, запас + доставлено − потери), т", f"=MIN(B{r_dem},B{r_open}+C{{deliv}}-C{{loss}})", row["served_total_t"]),
        ("Дефицит = спрос − выдано, т", f"=B{r_dem}-C{{served}}", row["shortage_total_t"]),
        ("Запас на конец = начало + доставлено − потери − выдано, т", f"=B{r_open}+C{{deliv}}-C{{loss}}-C{{served}}", row["closing_inventory_t"]),
        ("SL общий = выдано / спрос", f"=C{{served}}/B{r_dem}", row["SL_total"]),
        ("Закупка всего, млн", f"=SUM(I{first}:I{last})", row["procurement_mln"]),
        ("Резервирование всего, млн", f"=SUM(J{first}:J{last})", row["reservation_mln"]),
        ("Хранение ≈ ставка × средний запас (ядро: по дням), млн", f"=B{r_hold}*(B{r_open}+C{{close}})/2", row["holding_mln"]),
        ("OPEX ZBO, млн", f"=B{r_opex}", row["opex_zbo_mln"]),
        ("CAPEX года, млн (из плана)", row["capex_mln"], row["capex_mln"]),
        ("Расходы года (с хранением по ядру), млн", f"=C{{proc}}+C{{resv}}+D{{hold}}+C{{opex}}+C{{capex}}", row["total_expense_mln"]),
        ("Коэффициент дисконта = 1/(1+r)^(год−t0)", f"=1/(1+B{r_rate})^({year}-B{r_t0})", row["discount_factor"]),
        ("PV расходов, млн", f"=C{{total}}*C{{df}}", row["discounted_expense_mln"]),
        ("Стоимость тонны served, млн/т", f"=C{{total}}/C{{served}}", row["cost_per_served_t_mln"]),
    ]
    tags = ["ordered", "deliv", "loss", "reserve", "rescheck", "served", "short", "close", "sl", "proc", "resv", "hold", "opex", "capex", "total", "df", "pv", "cpt"]
    rows_idx = {}
    start = ws.max_row + 1
    for i, tag in enumerate(tags):
        rows_idx[tag] = start + i
    for (name, formula, core), tag in zip(calc, tags):
        rr = rows_idx[tag]
        f = formula
        if isinstance(f, str):
            for t, ri in rows_idx.items():
                f = f.replace("{" + t + "}", str(ri))
        ws.append([name, "", f, core, f"=IF(ISNUMBER(C{rr}),C{rr}-D{rr},\"\")"])
    ws.append([])
    ws.append(["Примечание: хранение в ядре считается по средневзвешенному дневному запасу; в контрольном примере приведена оценка по полусумме начала и конца года, поэтому для расходов года используется значение ядра. Остальные строки должны совпадать с точностью округления."])
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["C"].width = 44
    for c in "BDEFGHIJKL":
        ws.column_dimensions[c].width = 16


def main() -> int:
    case = load_case(CASE_DIR)
    params, _ = load_assumptions()
    scen = all_scenarios(case)
    wb = Workbook()
    wb.remove(wb.active)
    sheet_for_year(wb, "2037_BASE", case, Plan.load(CONFIG_DIR / "plans" / f"{CHOSEN}.json"), scen["BASE"], params, 2037)
    sheet_for_year(wb, "2039_STRESS_prepared", case, Plan.load(CONFIG_DIR / "plans" / f"{CHOSEN}_stress_prepared.json"), scen["MANDATORY_STRESS"], params, 2039)
    sheet_for_year(wb, "2039_STRESS_fixed", case, Plan.load(CONFIG_DIR / "plans" / f"{CHOSEN}.json"), scen["MANDATORY_STRESS"], params, 2039)
    out = ROOT / "docs" / "control_example.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"written {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
